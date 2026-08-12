"""
Processes an Excel workbook and a YAML file to update indicator information based
on matching principles, outcomes, and codes within rows of the spreadsheet.

This script reads an Excel file containing principles, outcomes, and indicators,
parses the data row by row, and updates the corresponding YAML structure. It ensures
that each indicator is categorized into "not-achieved," "partially-achieved," or
"achieved" sections based on its location in the spreadsheet. The updated YAML
is then written back to disk.

This script is designed to be used in conjunction with the generate_excel.py script,
which generates an Excel workbook containing principles, outcomes, and indicator data.
The govassure team is expected to update the Excel workbook and send it back to us,
at which point this script will process the updated workbook and update the corresponding
YAML structure.

Attributes:
    yaml (ruamel.yaml.YAML): Instance of the YAML processor, configured for anchor
        and alias preservation.
    wb (openpyxl.Workbook): Loaded Excel workbook containing principles, outcomes,
        and indicator data.
    data (dict | None): The parsed YAML structure to be updated. Initially, this is
        read from the input YAML file.

Raises:
    FileNotFoundError: Raised if the input files (Excel or YAML) do not exist.
    KeyError: Raised if keys in the YAML file or Excel workbook are missing or not
        in the expected format.
"""

import re

from openpyxl import load_workbook
from ruamel.yaml import YAML
from typing_extensions import Any

# Using this to load yaml as it preserves anchors and aliases
yaml = YAML()
yaml.width = 90
wb = load_workbook("cyber-assessment-framework-v4.0-import.xlsx")
data: dict[str, Any] | None
with open("cyber-assessment-framework-v4.0.yaml", "r") as f:
    data = yaml.load(f)
    for ws in wb.worksheets:
        title = re.sub(r"\s+", "-", ws.title)
        adding_indicators = False
        principle: str | None = None
        outcome: str | None = None

        print(f"Processing {title}")
        for row in ws.iter_rows(min_row=0, values_only=True):
            if row[1] and row[1].startswith("Principle"):
                match = re.match(r"Principle\s([A-Z]\d)-.*", str(row[1]))
                if match and match.groups():
                    principle = match.group(1)
                    print(principle)
                    continue
            if row[1] and row[1].startswith("Outcome"):
                re_match = re.match(r"Outcome\s([A-Z]\d\.[a-z])-.*", str(row[1]))
                if re_match and re_match.groups():
                    outcome = re_match.group(1)
                    print(outcome)
                    continue
            if row[1] and row[1].startswith("Not achieved Code"):
                adding_indicators = True
                continue
            # Check all the keys
            if not row[1] and not row[4] and not row[7]:
                adding_indicators = False
                continue
            if adding_indicators:
                indicator_mapping = {1: "not-achieved", 4: "partially-achieved", 7: "achieved"}
                # Process indicators by row, Not achieved, Partially achieved, Achieved
                for idx in range(1, 8, 3):
                    if row[idx] and re.match(rf"{outcome}\.\d+", str(row[idx])):
                        code = row[idx]
                        description = row[idx + 1]
                        print(code, description)
                        if code:
                            data_outcome = data["objectives"][principle[0]]["principles"][principle]["outcomes"][outcome]  # type: ignore[index]
                            indicator_category = indicator_mapping[idx]
                            data_outcome["indicators"][indicator_category][code][
                                "description"
                            ] = description  # type: ignore[index]

with open("cyber-assessment-framework-v4.0-updated.yaml", "w") as f:
    yaml.dump(data, f)
